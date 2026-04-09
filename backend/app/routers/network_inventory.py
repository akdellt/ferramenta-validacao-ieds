from fastapi import APIRouter, UploadFile, File, Depends, HTTPException
from sqlalchemy.orm import Session
from app.database import get_db
from app.models import NetworkIED
from openpyxl import load_workbook
from io import BytesIO

router = APIRouter(
    prefix="/inventory",
    tags=["Inventory"]
)

@router.post("/import")
async def import_ieds(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename or not file.filename.lower().endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Formato inválido. Use apenas .xlsx")

    content = await file.read()
    
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
        sheet = wb.active

        if not sheet:
            raise ValueError("Não foi possível ler a planilha ativa.")

        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        header = {str(cell).strip().lower(): i for i, cell in enumerate(header_row) if cell}
        
        required = {'nome', 'subestação', 'componente', 'modelo', 'ip'}
        if not required.issubset(header.keys()):
            missing = required - set(header.keys())
            raise ValueError(f"Planilha fora do padrão. Colunas ausentes: {', '.join(missing)}")
        
        db.query(NetworkIED).delete()

        rows_added = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if all(cell is None for cell in row):
                continue

            for key in required:
                val = row[header[key]]
                if val is None or str(val).strip() == "":
                    raise ValueError(f"Campo obrigatório '{key}' está vazio para o IED na linha {rows_added + 2}.")

            tipo_idx = header.get('tipo')
            tipo_val = str(row[tipo_idx]).strip().upper() if (tipo_idx is not None and row[tipo_idx]) else "TELNET"

            new_ied = NetworkIED(
                name=str(row[header['nome']]).strip().upper(),
                substation=str(row[header['subestação']]).strip().upper(),
                component_name=str(row[header['componente']]).strip().upper(),
                relay_model=str(row[header['modelo']]).strip().upper(),
                ip_address=str(row[header['ip']]).strip(),
                connection_type=tipo_val
            )
            db.add(new_ied)
            rows_added += 1

        db.commit()
        return {"message": f"{rows_added} IEDs carregados com sucesso"}

    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"{str(e)}")